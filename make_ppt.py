from pptx import Presentation
from pptx.util import Inches, Pt
prs = Presentation(); prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)
def add(title, bullets):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    tf = s.shapes.placeholders[1].text_frame
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b; p.font.size = Pt(20); p.space_after = Pt(10)
    return s
add("ShopCluster — Customer Segmentation with K-Means", ["By Vikranth Reddy", "B.Tech CSE (AI & ML), Woxsen University", "Daily AI/ML Project — 27 October 2026"])
add("The Problem", ["Shops treat every customer the same — same ads, same offers.", "A student and a high earner do not want the same deal.", "Without grouping, marketing money is wasted on the wrong people.", "Manual grouping by gut feeling does not scale to thousands of customers."])
add("The Solution", ["ShopCluster finds natural customer groups automatically.", "Enter age, income and spending score — the demo names the segment.", "Each segment gets a tailored offer: drops, premium perks, bundles, or honest support.", "No labels needed: this is unsupervised learning."])
add("Tech Stack", ["Python + scikit-learn for training (K-Means, k = 4).", "300 synthetic shoppers, seed 42 — reproducible.", "StandardScaler-style standardisation before clustering.", "Single-page HTML + JavaScript demo — no backend, runs in the browser.", "GitHub for code, Vercel for hosting."])
add("Architecture", ["1. Generate synthetic shoppers around 4 archetypes.", "2. Standardise age, income and spending score.", "3. K-Means: move 4 centres until shoppers stop switching groups.", "4. Export centres, means and standard deviations to model.json.", "5. Browser demo loads the same centres and assigns the nearest segment."])
add("Live Demo", ["Try it at shopcluster.vercel.app.", "Example 1: age 24, income $45k, spending 82 → Young Spenders.", "Example 2: age 48, income $92k, spending 68 → Affluent Regulars.", "The page shows the distance to every centre, so the answer is explainable."])
add("Results", ["Silhouette score: 0.53 — well-separated, sensible clusters.", "Segment sizes: 76, 76, 76, 72 — balanced, no empty cluster.", "Centres match the archetypes: e.g. Young Spenders ≈ age 25, $43k, spending 79.", "JavaScript demo reproduces Python assignments exactly on all test cases."])
add("Future Scope", ["Add real store data: purchase history, visit frequency, recency.", "Try more features and compare k = 3 to 6 with elbow + silhouette plots.", "Add a scatter plot of the clusters in the demo.", "Recommend a product per segment, not just an offer type."])
add("Thank You", ["ShopCluster — find the group, fit the offer.", "Code: github.com/HVRGAMING2705/shopcluster", "Questions welcome."])
prs.save("/home/hatch/workspace/goals/daily-ai-ml-project-factory/files/AIML_ShopCluster_2026-10-27.pptx")
print("ppt saved")
